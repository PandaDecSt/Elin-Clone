#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
build_sources.py  —  Elin SourceData 烘焙器 (faithful port of Elin's xlsx reading)

Elin 的 Source 体系约定 (SourceBlock/SourceCard/SourceChara/Lang):
  - 第 1 行 = 列名 (header)
  - 第 2 行 = 类型标记 (type row): int / float / bool / string / color / elements / int[] / string[] / alias ...
  - 第 3 行起 = 数据行
  - 第 1 列通常是主键 `id` (int 或 string)
  - `alias` 列: 该行是某 id 的别名数据, 合并进目标行 (默认填充)

输出: data/elin_source/sources.json
  {
    "SourceBlock":  { "<Sheet>": { "rows":[...], "byId":{ "id": row } }, ... },
    "SourceCard":   { ... },
    "SourceChara":  { ... },
    "Lang":        { "General":{id:{filter,jp,en}}, "Game":{...}, "Note":{...}, "List":{...}, "Word":{id:{group,jp,en}} }
  }

中文策略: 源数据只有 name_JP(日文) / name(英文) / Lang.text(英文). 本脚本不臆造中文,
运行时通过 lang_zh.json 覆盖 (见 build_lang_zh.py). 解析层把 name_JP/name 都保留.
"""
import openpyxl, os, json, re

SRC = 'data/sources'
OUT = 'data/elin_source/sources.json'

# ---------------------------------------------------------------- type coercion
def coerce(raw, ttype):
    """Convert a raw cell value according to Elin type marker."""
    if raw is None:
        return None
    t = (ttype or '').strip()
    # array types
    if t in ('int[]', 'string[]', 'elements', 'float[]') or t.endswith('[]'):
        if isinstance(raw, str):
            parts = re.split(r'[,|]', raw)
            items = [p.strip() for p in parts if p.strip() != '']
            if t.startswith('int'):
                out = []
                for p in items:
                    try: out.append(int(p))
                    except: out.append(p)
                return out
            return items
        if isinstance(raw, (list, tuple)):
            return list(raw)
        return [raw]
    if t == 'bool':
        if isinstance(raw, bool): return raw
        if isinstance(raw, (int, float)): return raw != 0
        return str(raw).strip().lower() in ('true', '1', 'yes')
    if t == 'int':
        if isinstance(raw, bool): return int(raw)
        if isinstance(raw, (int, float)): return int(raw)
        if isinstance(raw, str):
            s = raw.strip()
            if s == '': return None
            try: return int(float(s))
            except: return s
        return raw
    if t in ('float', 'double'):
        if isinstance(raw, (int, float)): return float(raw)
        if isinstance(raw, str):
            s = raw.strip()
            if s == '': return None
            try: return float(s)
            except: return s
        return raw
    # string / color / alias / unknown -> trust openpyxl native type, but normalize
    if isinstance(raw, str):
        s = raw.strip()
        return s if s != '' else None
    return raw


def parse_sheet(ws):
    """Return (rows_list, byId_dict, col_types) for a generic Source sheet."""
    all_rows = list(ws.iter_rows(values_only=True))
    if len(all_rows) < 2:
        return [], {}, {}
    header = all_rows[0]
    types = all_rows[1] if len(all_rows) > 1 else []
    # map column index -> name (skip None headers)
    cols = {}
    for i, h in enumerate(header):
        if h is not None and str(h).strip() != '':
            cols[i] = str(h).strip()
    col_types = {cols[i]: (types[i] if i < len(types) else None) for i in cols}
    id_col = 0 if (0 in cols and cols[0].lower() == 'id') else None
    rows = []
    byId = {}
    for r in all_rows[2:]:
        row = {}
        for i, name in cols.items():
            raw = r[i] if i < len(r) else None
            row[name] = coerce(raw, col_types[name])
        # skip fully-empty rows
        if not any(v is not None for v in row.values()):
            continue
        rows.append(row)
        if id_col is not None:
            key = row.get(cols[id_col])
            if key is not None and str(key).strip() != '':
                byId[str(key)] = row
    return rows, byId, col_types


def parse_lang_sheet(ws, sheet_name):
    """Lang sheets carry localized text, not entity data."""
    all_rows = list(ws.iter_rows(values_only=True))
    if len(all_rows) < 2:
        return {}
    header = all_rows[0]
    types = all_rows[1] if len(all_rows) > 1 else []
    def tcol(name):
        for i, h in enumerate(header):
            if h == name: return i, (types[i] if i < len(types) else None)
        return None, None
    out = {}
    if sheet_name == 'Word':
        i_id, _ = tcol('id'); i_group, _ = tcol('group'); i_jp, _ = tcol('name_JP'); i_en, _ = tcol('name')
        for r in all_rows[2:]:
            if i_id is None or r[i_id] is None: continue
            out[str(r[i_id])] = {
                'group': r[i_group] if i_group is not None else None,
                'jp': r[i_jp] if i_jp is not None else None,
                'en': r[i_en] if i_en is not None else None,
            }
    elif sheet_name == 'Note':
        i_id, _ = tcol('id'); i_jp, _ = tcol('text_JP'); i_en, _ = tcol('text')
        for r in all_rows[2:]:
            if i_id is None or r[i_id] is None: continue
            out[str(r[i_id])] = {'jp': r[i_jp], 'en': r[i_en]}
    elif sheet_name == 'List':
        i_id, _ = tcol('id'); i_filter, _ = tcol('filter'); i_jp, tjp = tcol('text_JP'); i_en, ten = tcol('text')
        for r in all_rows[2:]:
            if i_id is None or r[i_id] is None: continue
            out[str(r[i_id])] = {
                'filter': r[i_filter],
                'jp': coerce(r[i_jp], tjp),
                'en': coerce(r[i_en], ten),
            }
    else:  # General / Game
        i_id, _ = tcol('id'); i_filter, _ = tcol('filter')
        i_en, ten = tcol('text'); i_jp, tjp = tcol('text_JP')
        extra = {}
        if sheet_name == 'Game':
            for f in ('group', 'color', 'logColor', 'sound', 'effect'):
                extra[f] = tcol(f)[0]
        for r in all_rows[2:]:
            if i_id is None or r[i_id] is None: continue
            rec = {'filter': r[i_filter] if i_filter is not None else None,
                   'jp': coerce(r[i_jp], tjp), 'en': coerce(r[i_en], ten)}
            for f, ci in extra.items():
                if ci is not None: rec[f] = r[ci]
            out[str(r[i_id])] = rec
    return out


def alias_merge(sheet_obj):
    """Elin alias: a row whose `alias` field points to another id gets merged
    into the target row (fills null fields of the target)."""
    rows = sheet_obj['rows']
    byId = sheet_obj['byId']
    if 'alias' not in (sheet_obj.get('colTypes') or {}):
        return
    for row in rows:
        al = row.get('alias')
        if al is None:
            continue
        target = byId.get(str(al))
        if target is None:
            continue
        for k, v in row.items():
            if v is None:
                continue
            if k == 'alias':
                continue
            if target.get(k) is None:
                target[k] = v


# ---------------------------------------------------------------- main
def main():
    result = {}
    # entity sources
    for fn in ['SourceBlock.xlsx', 'SourceCard.xlsx', 'SourceChara.xlsx']:
        base = fn.replace('.xlsx', '')
        wb = openpyxl.load_workbook(os.path.join(SRC, fn), read_only=True, data_only=True)
        container = {}
        for sn in wb.sheetnames:
            ws = wb[sn]
            rows, byId, colTypes = parse_sheet(ws)
            sheet_obj = {'rows': rows, 'byId': byId, 'colTypes': colTypes, 'count': len(rows)}
            alias_merge(sheet_obj)
            container[sn] = sheet_obj
            print(f'  [{base}/{sn}] {len(rows)} rows, byId={len(byId)}')
        wb.close()
        result[base] = container

    # Lang
    wb = openpyxl.load_workbook(os.path.join(SRC, 'Lang.xlsx'), read_only=True, data_only=True)
    lang = {}
    for sn in wb.sheetnames:
        lang[sn] = parse_lang_sheet(wb[sn], sn)
        print(f'  [Lang/{sn}] {len(lang[sn])} entries')
    wb.close()
    result['Lang'] = lang

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    json.dump(result, open(OUT, 'w'), ensure_ascii=False, indent=1)
    print(f'\nWrote {OUT}')

    # summary
    tot = 0
    for base, cont in result.items():
        if base == 'Lang':
            continue
        for sn, so in cont.items():
            tot += so['count']
    print(f'Total entity rows: {tot}')
    print(f'Lang entries: {sum(len(v) for v in result["Lang"].values())}')


if __name__ == '__main__':
    main()
